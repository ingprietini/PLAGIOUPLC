import io
import re
from pathlib import Path

SUPPORTED_EXTENSIONS = {
    ".txt": "text/plain",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".doc": "application/msword",
    ".odt": "application/vnd.oasis.opendocument.text",
}

MAX_FILE_SIZE = 50 * 1024 * 1024


def extract_text_from_bytes(content: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()

    if ext == ".txt":
        return content.decode("utf-8", errors="replace")

    if ext == ".pdf":
        return _extract_pdf(content)

    if ext in (".docx", ".doc"):
        return _extract_docx(content)

    if ext == ".odt":
        return _extract_odt(content)

    raise ValueError(f"Formato no soportado: {ext}")


def extract_text_from_file(filepath: str) -> str:
    with open(filepath, "rb") as f:
        content = f.read()
    return extract_text_from_bytes(content, filepath)


def _extract_pdf(content: bytes) -> str:
    import pdfplumber

    text_parts = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n\n".join(text_parts)


def _extract_docx(content: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(content))
    text_parts = [p.text for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                text_parts.append(" | ".join(cells))

    return "\n".join(text_parts)


def _extract_odt(content: bytes) -> str:
    import zipfile
    import xml.etree.ElementTree as ET

    with zipfile.ZipFile(io.BytesIO(content)) as z:
        if "content.xml" not in z.namelist():
            raise ValueError("ODT inválido: no se encontró content.xml")

        xml_content = z.read("content.xml")
        root = ET.fromstring(xml_content)

        ns = {
            "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
            "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
        }

        text_parts = []
        for p in root.iter("{urn:oasis:names:tc:opendocument:xmlns:text:1.0}p"):
            text = "".join(p.itertext()).strip()
            if text:
                text_parts.append(text)

        return "\n".join(text_parts)


def get_word_count(text: str) -> int:
    return len(text.split())


def truncate_text(text: str, max_words: int = 100000) -> str:
    words = text.split()
    if len(words) <= max_words:
        return text
    return " ".join(words[:max_words])
