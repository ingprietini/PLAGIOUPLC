import json
import asyncio
import logging
import io
import re
import importlib
import html
from datetime import datetime
from contextlib import asynccontextmanager
from fastapi import FastAPI, WebSocket, WebSocketDisconnect, Form, UploadFile, File, HTTPException, Depends
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from sqlalchemy.orm import Session
from docx import Document

from config import MAX_WORD_LIMIT, MAX_UPLOAD_SIZE, SIMILARITY_THRESHOLD, MAX_SOURCES_TO_CHECK
from models.database import init_db, dispose_engine, get_db, Analysis, Source, Suggestion
from parser import extract_text_from_bytes, get_word_count, truncate_text, SUPPORTED_EXTENSIONS
from plagiarism.winnowing import analyze_plagiarism
from plagiarism.tfidf import compute_similarity
from plagiarism.search import search_chunks, hydrate_sources_content, close_client
from plagiarism.source_normalizer import normalize_sources
from plagiarism.scoring import rank_sources, build_stats
from nlp.readability import analyze_readability
from nlp.suggestions import get_improvement_suggestions
from apa.formatter import format_apa, format_mla, format_chicago, CITATION_STYLES
from ai_detector.detector import detect_ai_text

logger = logging.getLogger("antiplagio")


def _parse_source_payload(matched_text: str | None) -> tuple[list, dict]:
    if not matched_text:
        return [], {}
    try:
        data = json.loads(matched_text)
    except Exception:
        return [], {}

    if isinstance(data, list):
        return data, {}
    if isinstance(data, dict):
        return data.get("segments", []), data.get("metrics", {})
    return [], {}


def _safe_filename(value: str) -> str:
    compact = re.sub(r"\s+", "_", (value or "informe").strip())
    compact = re.sub(r"[^a-zA-Z0-9_\-]", "", compact)
    return compact[:80] or "informe"


def _segment_excerpt(segment) -> str:
    if isinstance(segment, str):
        return segment.strip()
    if isinstance(segment, dict):
        source_excerpt = (segment.get("source_excerpt") or "").strip()
        text_excerpt = (segment.get("text_excerpt") or "").strip()
        phrase = (segment.get("matched_phrase") or "").strip()
        parts = []
        if phrase:
            parts.append(f"Frase: {phrase}")
        if source_excerpt:
            parts.append(f"Web: {source_excerpt}")
        if text_excerpt:
            parts.append(f"Documento: {text_excerpt}")
        return " | ".join(parts)
    return ""


def _iter_text_blocks(text: str, max_chars: int = 1400) -> list[str]:
    value = (text or "").replace("\r\n", "\n")
    blocks: list[str] = []
    for line in value.split("\n"):
        clean = line.strip()
        if not clean:
            continue
        while len(clean) > max_chars:
            blocks.append(clean[:max_chars])
            clean = clean[max_chars:]
        blocks.append(clean)
    return blocks or [""]


def _collect_plagiarism_notes(sources: list) -> list[dict]:
    notes: list[dict] = []
    seen_phrases: set[str] = set()
    note_id = 1

    for source in sources:
        segments, _ = _parse_source_payload(source.matched_text)
        for segment in segments:
            if not isinstance(segment, dict):
                continue

            phrase = (segment.get("matched_phrase") or "").strip()
            text_excerpt = (segment.get("text_excerpt") or "").strip()
            source_excerpt = (segment.get("source_excerpt") or "").strip()
            candidate = phrase or text_excerpt

            if not candidate:
                continue

            normalized = " ".join(candidate.lower().split())
            if normalized in seen_phrases:
                continue
            seen_phrases.add(normalized)

            token = f"__PLG_{note_id}__"
            notes.append(
                {
                    "id": note_id,
                    "token": token,
                    "phrase": candidate,
                    "text_excerpt": text_excerpt,
                    "source_excerpt": source_excerpt,
                    "source_title": source.title or "Sin título",
                    "source_url": source.url,
                    "similarity": source.similarity,
                    "hits": 0,
                }
            )
            note_id += 1

    return notes


def _annotate_text_block(block: str, notes: list[dict]) -> str:
    annotated = block

    # Replace longer phrases first to avoid partial replacements on shorter phrases.
    ordered_notes = sorted(notes, key=lambda n: len(n["phrase"]), reverse=True)

    for note in ordered_notes:
        phrase = note["phrase"]
        if len(phrase.split()) < 4:
            continue

        pattern = re.compile(re.escape(phrase), re.IGNORECASE)
        annotated, count = pattern.subn(note["token"], annotated)
        if count > 0:
            note["hits"] += count

    escaped = html.escape(annotated)
    for note in ordered_notes:
        token = note["token"]
        if token not in escaped:
            continue

        phrase_markup = html.escape(note["phrase"])
        replacement = (
            f'<font backColor="#FFE08A" color="#8A1C1C"><b>{phrase_markup}</b></font>'
            f'<super><font color="#8A1C1C">[{note["id"]}]</font></super>'
        )
        escaped = escaped.replace(token, replacement)

    return escaped


@asynccontextmanager
async def lifespan(app: FastAPI):
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    init_db()
    yield
    dispose_engine()
    await close_client()


app = FastAPI(title="AntiPlagio API", version="1.0.0", lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.exception_handler(HTTPException)
async def http_exception_handler(request, exc: HTTPException):
    return JSONResponse(
        status_code=exc.status_code,
        content={"status": "error", "detail": exc.detail},
    )


@app.exception_handler(Exception)
async def general_exception_handler(request, exc: Exception):
    return JSONResponse(
        status_code=500,
        content={"status": "error", "detail": "Error interno del servidor"},
    )


@app.get("/api/health")
async def health():
    return {"status": "ok", "timestamp": datetime.utcnow().isoformat()}


@app.get("/")
async def root():
    return {
        "status": "ok",
        "service": "AntiPlagio API",
        "docs": "/docs",
        "health": "/api/health",
    }


ALLOWED_EXTENSIONS = {".txt", ".pdf", ".docx", ".doc", ".odt"}


@app.post("/api/analyze")
async def create_analysis(text: str = Form(...), title: str = Form(None), db: Session = Depends(get_db)):
    word_count = len(text.split())
    if word_count > MAX_WORD_LIMIT:
        raise HTTPException(status_code=400, detail=f"El texto excede el límite de {MAX_WORD_LIMIT} palabras")
    if word_count < 50:
        raise HTTPException(status_code=400, detail="El texto debe tener al menos 50 palabras")

    analysis = Analysis(
        title=title or f"Análisis {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        original_text=text,
        word_count=word_count,
        status="pending",
    )
    db.add(analysis)
    db.commit()

    return {"status": "ok", "id": analysis.id, "word_count": word_count}


@app.post("/api/upload")
async def upload_analysis(file: UploadFile = File(...), db: Session = Depends(get_db)):
    ext = "." + file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"Formato no soportado: {ext}. Permitidos: {', '.join(sorted(ALLOWED_EXTENSIONS))}",
        )

    content = await file.read()
    if len(content) > MAX_UPLOAD_SIZE:
        raise HTTPException(status_code=400, detail="El archivo excede el tamaño máximo de 50 MB")

    try:
        text = extract_text_from_bytes(content, file.filename)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error al leer el archivo: {str(e)}")

    if not text.strip():
        raise HTTPException(status_code=400, detail="No se pudo extraer texto del archivo")

    word_count = get_word_count(text)
    if word_count > MAX_WORD_LIMIT:
        text = truncate_text(text, MAX_WORD_LIMIT)
        word_count = MAX_WORD_LIMIT

    title = file.filename.rsplit(".", 1)[0].replace("_", " ").replace("-", " ").title()

    analysis = Analysis(
        title=title,
        original_text=text,
        word_count=word_count,
        status="pending",
    )
    db.add(analysis)
    db.commit()

    return {"status": "ok", "id": analysis.id, "word_count": word_count, "filename": file.filename}


@app.get("/api/analyze/{analysis_id}")
async def get_analysis(analysis_id: str, db: Session = Depends(get_db)):
    analysis = db.query(Analysis).filter(Analysis.id == analysis_id).first()
    if not analysis:
        raise HTTPException(status_code=404, detail="Análisis no encontrado")

    sources = db.query(Source).filter(Source.analysis_id == analysis_id).all()
    suggestions = db.query(Suggestion).filter(Suggestion.analysis_id == analysis_id).all()

    source_items = []
    for s in sources:
        segments, metrics = _parse_source_payload(s.matched_text)
        source_items.append(
            {
                "id": s.id,
                "url": s.url,
                "title": s.title,
                "similarity": s.similarity,
                "matched_text": segments,
                "source_snippet": s.source_snippet,
                "match_confidence": metrics.get("match_confidence"),
                "content_quality": metrics.get("content_quality"),
                "short_page_penalty": metrics.get("short_page_penalty"),
                "source_word_count": metrics.get("source_word_count"),
            }
        )

    return {
        "id": analysis.id,
        "title": analysis.title,
        "word_count": analysis.word_count,
        "status": analysis.status,
        "progress": analysis.progress,
        "error": analysis.error,
        "stats": analysis.stats,
        "sources": source_items,
        "suggestions": [
            {
                "id": sg.id,
                "type": sg.type,
                "severity": sg.severity,
                "suggestion": sg.suggestion_text,
                "explanation": sg.explanation,
                "improved_example": sg.improved_example,
                "original_text": sg.original_text,
                "position_start": sg.position_start,
                "position_end": sg.position_end,
            }
            for sg in suggestions
        ],
        "created_at": analysis.created_at.isoformat() if analysis.created_at else None,
        "completed_at": analysis.completed_at.isoformat() if analysis.completed_at else None,
    }


@app.get("/api/analyze/{analysis_id}/report/docx")
async def download_analysis_report_docx(analysis_id: str, db: Session = Depends(get_db)):
    analysis = db.query(Analysis).filter(Analysis.id == analysis_id).first()
    if not analysis:
        raise HTTPException(status_code=404, detail="Análisis no encontrado")

    sources = db.query(Source).filter(Source.analysis_id == analysis_id).all()
    suggestions = db.query(Suggestion).filter(Suggestion.analysis_id == analysis_id).all()
    stats = analysis.stats or {}

    doc = Document()

    # Primera hoja: resumen consolidado del análisis.
    doc.add_heading("Informe de Detección de Plagio", level=0)
    doc.add_paragraph(f"Análisis: {analysis.title}")
    doc.add_paragraph(f"ID: {analysis.id}")
    doc.add_paragraph(f"Fecha: {analysis.completed_at.isoformat() if analysis.completed_at else datetime.utcnow().isoformat()}")

    doc.add_heading("Estadísticas principales", level=1)
    stats_lines = [
        f"Palabras analizadas: {analysis.word_count}",
        f"Fuentes consultadas: {stats.get('sources_checked', 0)}",
        f"Fuentes detectadas: {len(sources)}",
        f"Similitud global: {stats.get('overall_similarity', 0)}%",
        f"Similitud máxima: {stats.get('max_similarity', 0)}%",
        f"Similitud promedio (sobre umbral): {stats.get('avg_similarity', 0)}%",
        f"Fuentes sobre umbral: {stats.get('sources_above_threshold', 0)}",
    ]
    for line in stats_lines:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Fuentes encontradas", level=1)
    if not sources:
        doc.add_paragraph("No se encontraron fuentes con coincidencias.")
    else:
        for idx, source in enumerate(sources, start=1):
            doc.add_paragraph(f"{idx}. {source.title or 'Sin título'} ({source.similarity}%)", style="List Number")
            doc.add_paragraph(f"URL: {source.url}")

    doc.add_heading("Hallazgos y sugerencias", level=1)
    if not suggestions:
        doc.add_paragraph("No se registraron sugerencias.")
    else:
        for idx, suggestion in enumerate(suggestions, start=1):
            doc.add_paragraph(
                f"{idx}. [{suggestion.severity}] {suggestion.type}: {suggestion.suggestion_text}",
                style="List Number",
            )
            if suggestion.explanation:
                doc.add_paragraph(f"Detalle: {suggestion.explanation}")

    doc.add_page_break()

    # Páginas siguientes: detalle de coincidencias por fuente.
    doc.add_heading("Detalle de coincidencias por fuente", level=1)
    if not sources:
        doc.add_paragraph("Sin detalle de fuentes.")
    else:
        for idx, source in enumerate(sources, start=1):
            doc.add_heading(f"Fuente {idx}: {source.title or 'Sin título'}", level=2)
            doc.add_paragraph(f"URL: {source.url}")
            doc.add_paragraph(f"Similitud: {source.similarity}%")

            segments, metrics = _parse_source_payload(source.matched_text)
            if metrics:
                doc.add_paragraph(
                    "Métricas: "
                    f"confianza={metrics.get('match_confidence', 0)}, "
                    f"calidad={metrics.get('content_quality', 0)}, "
                    f"penalización_corta={metrics.get('short_page_penalty', 0)}"
                )

            if not segments:
                doc.add_paragraph("Sin cita textual verificable en esta fuente.")
            else:
                for seg in segments[:15]:
                    excerpt = _segment_excerpt(seg)
                    if excerpt:
                        doc.add_paragraph(excerpt, style="List Bullet")

    doc.add_page_break()
    doc.add_heading("Texto original cargado", level=1)
    for block in _iter_text_blocks(analysis.original_text, max_chars=1200):
        doc.add_paragraph(block)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    filename = _safe_filename(analysis.title)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}_reporte.docx"'},
    )


@app.get("/api/analyze/{analysis_id}/report/pdf")
async def download_analysis_report_pdf(analysis_id: str, db: Session = Depends(get_db)):
    analysis = db.query(Analysis).filter(Analysis.id == analysis_id).first()
    if not analysis:
        raise HTTPException(status_code=404, detail="Análisis no encontrado")

    try:
        pagesizes = importlib.import_module("reportlab.lib.pagesizes")
        styles_mod = importlib.import_module("reportlab.lib.styles")
        units = importlib.import_module("reportlab.lib.units")
        platypus = importlib.import_module("reportlab.platypus")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"No se pudo generar PDF: {e}")

    A4 = pagesizes.A4
    ParagraphStyle = styles_mod.ParagraphStyle
    getSampleStyleSheet = styles_mod.getSampleStyleSheet
    cm = units.cm
    PageBreak = platypus.PageBreak
    Paragraph = platypus.Paragraph
    SimpleDocTemplate = platypus.SimpleDocTemplate
    Spacer = platypus.Spacer

    sources = db.query(Source).filter(Source.analysis_id == analysis_id).all()
    suggestions = db.query(Suggestion).filter(Suggestion.analysis_id == analysis_id).all()
    stats = analysis.stats or {}

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=1.8 * cm,
        rightMargin=1.8 * cm,
        topMargin=1.8 * cm,
        bottomMargin=1.8 * cm,
        title=f"Informe de Plagio - {analysis.title}",
    )

    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "BodySmall",
        parent=styles["BodyText"],
        fontSize=9,
        leading=12,
        spaceAfter=4,
    )
    comment_style = ParagraphStyle(
        "CommentBox",
        parent=body,
        backColor="#FFF8E1",
        borderColor="#D19A00",
        borderWidth=0.6,
        borderPadding=5,
        spaceAfter=7,
    )
    story = []

    notes = _collect_plagiarism_notes(sources)

    # Primera hoja: resumen completo.
    story.append(Paragraph("Informe de Detección de Plagio", styles["Title"]))
    story.append(Paragraph(f"Análisis: {analysis.title}", body))
    story.append(Paragraph(f"ID: {analysis.id}", body))
    story.append(Paragraph(
        f"Fecha: {analysis.completed_at.isoformat() if analysis.completed_at else datetime.utcnow().isoformat()}",
        body,
    ))
    story.append(Spacer(1, 8))

    story.append(Paragraph("Estadísticas principales", styles["Heading2"]))
    stats_lines = [
        f"Palabras analizadas: {analysis.word_count}",
        f"Fuentes consultadas: {stats.get('sources_checked', 0)}",
        f"Fuentes detectadas: {len(sources)}",
        f"Similitud global: {stats.get('overall_similarity', 0)}%",
        f"Similitud máxima: {stats.get('max_similarity', 0)}%",
        f"Similitud promedio (sobre umbral): {stats.get('avg_similarity', 0)}%",
        f"Fuentes sobre umbral: {stats.get('sources_above_threshold', 0)}",
    ]
    for line in stats_lines:
        story.append(Paragraph(f"- {line}", body))

    story.append(Spacer(1, 8))
    story.append(Paragraph("Fuentes encontradas", styles["Heading2"]))
    if not sources:
        story.append(Paragraph("No se encontraron fuentes con coincidencias.", body))
    else:
        for idx, source in enumerate(sources, start=1):
            story.append(Paragraph(f"{idx}. {source.title or 'Sin título'} ({source.similarity}%)", body))
            story.append(Paragraph(f"URL: {source.url}", body))

    story.append(Spacer(1, 8))
    story.append(Paragraph("Hallazgos y sugerencias", styles["Heading2"]))
    if not suggestions:
        story.append(Paragraph("No se registraron sugerencias.", body))
    else:
        for idx, suggestion in enumerate(suggestions, start=1):
            story.append(Paragraph(
                f"{idx}. [{suggestion.severity}] {suggestion.type}: {suggestion.suggestion_text}",
                body,
            ))
            if suggestion.explanation:
                story.append(Paragraph(f"Detalle: {suggestion.explanation}", body))

    story.append(PageBreak())
    story.append(Paragraph("Detalle de coincidencias por fuente", styles["Heading1"]))
    if not sources:
        story.append(Paragraph("Sin detalle de fuentes.", body))
    else:
        for idx, source in enumerate(sources, start=1):
            story.append(Paragraph(f"Fuente {idx}: {source.title or 'Sin título'}", styles["Heading3"]))
            story.append(Paragraph(f"URL: {source.url}", body))
            story.append(Paragraph(f"Similitud: {source.similarity}%", body))

            segments, metrics = _parse_source_payload(source.matched_text)
            if metrics:
                story.append(Paragraph(
                    "Métricas: "
                    f"confianza={metrics.get('match_confidence', 0)}, "
                    f"calidad={metrics.get('content_quality', 0)}, "
                    f"penalización_corta={metrics.get('short_page_penalty', 0)}",
                    body,
                ))

            if not segments:
                story.append(Paragraph("Sin cita textual verificable en esta fuente.", body))
            else:
                for seg in segments[:15]:
                    excerpt = _segment_excerpt(seg)
                    if excerpt:
                        story.append(Paragraph(f"- {excerpt}", body))
            story.append(Spacer(1, 6))

    story.append(PageBreak())
    story.append(Paragraph("Comentarios de Plagio (Notas)", styles["Heading1"]))
    active_notes = [n for n in notes if n["hits"] > 0]
    if not active_notes:
        story.append(Paragraph("No se generaron anotaciones de plagio con coincidencia literal suficiente.", body))
    else:
        for note in active_notes[:120]:
            note_text = (
                f"<b>Nota [{note['id']}]</b> - Similitud fuente: {note['similarity']}%<br/>"
                f"<b>Fuente:</b> {html.escape(note['source_title'])}<br/>"
                f"<b>URL:</b> {html.escape(note['source_url'] or '')}<br/>"
                f"<b>Texto detectado:</b> {html.escape(note['phrase'])}<br/>"
                f"<b>Sección:</b> Texto original resaltado (final del informe)."
            )
            story.append(Paragraph(note_text, comment_style))

    story.append(PageBreak())
    story.append(Paragraph("Texto original cargado", styles["Heading1"]))
    for block in _iter_text_blocks(analysis.original_text, max_chars=1800):
        annotated_block = _annotate_text_block(block, notes)
        story.append(Paragraph(annotated_block, body))

    doc.build(story)
    buffer.seek(0)

    filename = _safe_filename(analysis.title)
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}_reporte.pdf"'},
    )


@app.websocket("/ws/analyze/{analysis_id}")
async def websocket_analyze(websocket: WebSocket, analysis_id: str):
    await websocket.accept()

    db = next(get_db())
    try:
        analysis = db.query(Analysis).filter(Analysis.id == analysis_id).first()
        if not analysis:
            await websocket.send_json({"type": "error", "message": "Análisis no encontrado"})
            await websocket.close()
            return

        analysis.status = "processing"
        analysis.progress = 0
        db.commit()

        await websocket.send_json({
            "type": "status",
            "status": "processing",
            "progress": 0,
            "message": "Iniciando análisis...",
        })

        text = analysis.original_text

        await websocket.send_json({
            "type": "progress",
            "progress": 5,
            "message": "Analizando legibilidad del texto...",
        })
        readability = analyze_readability(text)

        await websocket.send_json({
            "type": "progress",
            "progress": 8,
            "message": "Detectando texto generado por IA...",
        })
        ai_detection = detect_ai_text(text)

        await websocket.send_json({
            "type": "progress",
            "progress": 10,
            "message": "Buscando fuentes en la web (DuckDuckGo)...",
        })

        async def send_progress(msg: str):
            try:
                await websocket.send_json({"type": "progress", "progress": 10, "message": msg})
            except Exception:
                pass

        try:
            web_sources = await search_chunks(text, progress_callback=send_progress)
        except Exception:
            web_sources = []

        normalized_sources = normalize_sources(web_sources)

        await websocket.send_json({
            "type": "progress",
            "progress": 24,
            "message": "Verificando contenido real de páginas encontradas...",
        })

        try:
            normalized_sources = await hydrate_sources_content(
                normalized_sources,
                max_fetch=MAX_SOURCES_TO_CHECK,
                progress_callback=send_progress,
            )
        except Exception:
            pass

        normalized_sources = normalized_sources[:MAX_SOURCES_TO_CHECK]

        await websocket.send_json({
            "type": "progress",
            "progress": 30,
            "message": f"Se encontraron {len(normalized_sources)} posibles fuentes. Analizando similitudes...",
        })

        async def send_source_checking(source: dict, index: int, total: int):
            try:
                # Map source checking to 30-60% of global progress bar.
                source_progress = 30 + int((index / max(1, total)) * 30)
                await websocket.send_json(
                    {
                        "type": "source_checking",
                        "progress": source_progress,
                        "message": f"Analizando fuente {index}/{total}",
                        "source": {
                            "index": index,
                            "total": total,
                            "url": source.get("url", ""),
                            "title": source.get("title", ""),
                        },
                    }
                )
            except Exception:
                pass

        winnowing_results = await analyze_plagiarism(
            text,
            normalized_sources,
            progress_callback=send_source_checking,
        )

        await websocket.send_json({
            "type": "progress",
            "progress": 60,
            "message": "Calculando similitud TF-IDF...",
        })

        tfidf_scores = compute_similarity(text, normalized_sources)
        tfidf_map = {s["url"]: s["similarity"] for s in tfidf_scores}
        source_lookup = {s["url"]: s for s in normalized_sources}

        await websocket.send_json({
            "type": "progress",
            "progress": 75,
            "message": "Generando sugerencias de mejora...",
        })

        suggestions = get_improvement_suggestions(text)

        await websocket.send_json({
            "type": "progress",
            "progress": 85,
            "message": "Guardando resultados...",
        })

        sources_to_save = rank_sources(
            winnowing_results["results"],
            tfidf_map,
            source_lookup=source_lookup,
            limit=MAX_SOURCES_TO_CHECK,
        )

        for src in sources_to_save:
            payload = {
                "segments": src.get("matching_segments", []),
                "metrics": {
                    "raw_similarity": src.get("raw_similarity", src.get("similarity", 0)),
                    "match_confidence": src.get("match_confidence", 0),
                    "content_quality": src.get("content_quality", 0),
                    "short_page_penalty": src.get("short_page_penalty", 0),
                    "source_word_count": src.get("source_word_count", 0),
                },
            }

            source = Source(
                analysis_id=analysis_id,
                url=src["url"],
                title=src["title"],
                similarity=src["similarity"],
                matched_text=json.dumps(payload),
                source_snippet=source_lookup.get(src["url"], {}).get("snippet", "")[:500],
            )
            db.add(source)

        for sug in suggestions:
            db.add(
                Suggestion(
                    analysis_id=analysis_id,
                    type=sug["type"],
                    severity=sug["severity"],
                    suggestion_text=sug["suggestion"],
                    explanation=sug.get("explanation", ""),
                    improved_example=sug.get("improved_example"),
                )
            )

        stats = build_stats(
            sources=sources_to_save,
            readability=readability,
            ai_detection=ai_detection,
            word_count=analysis.word_count,
            similarity_threshold=SIMILARITY_THRESHOLD,
            sources_checked=len(normalized_sources),
        )

        analysis.stats = stats
        analysis.status = "completed"
        analysis.progress = 100
        analysis.completed_at = datetime.utcnow()
        db.commit()

        await websocket.send_json({
            "type": "progress",
            "progress": 100,
            "message": "Análisis completado",
        })

        await websocket.send_json({
            "type": "complete",
            "stats": stats,
            "sources_count": len(sources_to_save),
            "suggestions_count": len(suggestions),
        })

    except WebSocketDisconnect:
        try:
            db.rollback()
            analysis = db.query(Analysis).filter(Analysis.id == analysis_id).first()
            if analysis and analysis.status == "processing":
                analysis.status = "failed"
                analysis.error = "Conexión interrumpida"
                db.commit()
        except Exception as inner:
            logger.error("Error en WebSocketDisconnect handler: %s", inner)
    except Exception as e:
        logger.error("Error en WebSocket handler: %s", e)
        try:
            db.rollback()
            analysis = db.query(Analysis).filter(Analysis.id == analysis_id).first()
            if analysis:
                analysis.status = "failed"
                analysis.error = str(e)
                db.commit()
            await websocket.send_json({"type": "error", "message": str(e)})
        except Exception:
            pass
    finally:
        db.close()
        try:
            await websocket.close()
        except Exception:
            pass


@app.get("/api/styles")
async def get_styles():
    return {"status": "ok", "styles": CITATION_STYLES}


@app.post("/api/format/citation")
async def format_citation(
    style: str = Form(...),
    source_type: str = Form(...),
    author: str = Form(""),
    title: str = Form(""),
    year: str = Form(""),
    url: str = Form(""),
    publisher: str = Form(""),
    journal: str = Form(""),
    volume: str = Form(""),
    issue: str = Form(""),
    pages: str = Form(""),
    doi: str = Form(""),
    site_name: str = Form(""),
):
    formatters = {"apa": format_apa, "mla": format_mla, "chicago": format_chicago}
    formatter = formatters.get(style)
    if not formatter:
        raise HTTPException(status_code=400, detail=f"Estilo no soportado: {style}")

    data = {
        "author": author, "title": title, "year": year,
        "url": url, "publisher": publisher, "journal": journal,
        "volume": volume, "issue": issue, "pages": pages,
        "doi": doi, "site_name": site_name,
    }

    try:
        citation = formatter(source_type, data)
        return {"status": "ok", "citation": citation, "style": style}
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=8000)
